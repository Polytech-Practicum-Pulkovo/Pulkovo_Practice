import os
from pathlib import Path


def extract_text_from_pptx(pptx_path: str | os.PathLike) -> str:
    """Извлекает текст (включая таблицы) из PPTX постранично для передачи в LLM."""

    from pptx import Presentation

    path = Path(pptx_path)
    if not path.exists():
        raise FileNotFoundError(path)

    presentation = Presentation(str(path))
    slides_text = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip(" |"):
                        parts.append(row_text)
        if parts:
            slides_text.append(f"Слайд {index}:\n" + "\n".join(parts))

    return "\n\n".join(slides_text)
